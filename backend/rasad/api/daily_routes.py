"""Daily-use routes: alerts inbox, what changed since a date, search across the ledgers and
the document store, documents routed to a question, and brief export to Word."""

from __future__ import annotations

import io
import os
import re
from datetime import date, datetime, timedelta

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy import func, or_, select
from sqlalchemy.orm import Session

from ..auth import core as auth
from ..auth.models import User
from ..briefs import generate as briefs
from ..db import get_session
from ..engine import alerts as alert_engine
from ..models import Alert, Document, Forecast, PolicyObservation, Question, Source, ThreatEvent
from .deps import check_language, current_user, require_role

router = APIRouter(prefix="/api", tags=["daily"])

STOP = set("will the and for with from that this within days into over under than then about after before between during against among which their there these those have has been being would could should shall does into onto upon".split())


def _alert_row(a: Alert) -> dict:
    return {"id": a.id, "created_at": a.created_at.isoformat(timespec="seconds"), "kind": a.kind, "severity": a.severity, "question_code": a.question_code, "title_en": a.title_en, "title_ar": a.title_ar, "detail": a.detail, "acknowledged_by": a.acknowledged_by, "acknowledged_at": a.acknowledged_at.isoformat(timespec="seconds") if a.acknowledged_at else None}


@router.get("/alerts")
def list_alerts(all: bool = False, limit: int = 100, db: Session = Depends(get_session), user: User = Depends(current_user)) -> list[dict]:
    stmt = select(Alert).order_by(Alert.created_at.desc()).limit(limit)
    if not all:
        stmt = select(Alert).where(Alert.acknowledged_at.is_(None)).order_by(Alert.severity.desc(), Alert.created_at.desc()).limit(limit)
    return [_alert_row(a) for a in db.scalars(stmt)]


@router.post("/alerts/{alert_id}/ack")
def ack_alert(alert_id: int, db: Session = Depends(get_session), user: User = Depends(require_role("analyst"))) -> dict:
    a = db.get(Alert, alert_id)
    if a is None:
        raise HTTPException(404, "alert not found")
    if a.acknowledged_at is None:
        a.acknowledged_by, a.acknowledged_at = user.username, datetime.utcnow()
        db.commit()
        auth.audit(db, user, "alert_acknowledged", str(alert_id), {"kind": a.kind})
    return _alert_row(a)


@router.post("/alerts/ack-all")
def ack_all(db: Session = Depends(get_session), user: User = Depends(require_role("analyst"))) -> dict:
    n = 0
    for a in db.scalars(select(Alert).where(Alert.acknowledged_at.is_(None))):
        a.acknowledged_by, a.acknowledged_at = user.username, datetime.utcnow()
        n += 1
    db.commit()
    auth.audit(db, user, "alerts_acknowledged_all", None, {"n": n})
    return {"acknowledged": n}


@router.post("/alerts/scan")
def scan_alerts(db: Session = Depends(get_session), user: User = Depends(require_role("analyst"))) -> dict:
    """Run the poll-side alert checks now (the scheduler runs them daily)."""
    return {"raised": alert_engine.after_poll(db, since=datetime.utcnow() - timedelta(days=3))}


@router.get("/changes")
def changes(since: str | None = None, db: Session = Depends(get_session), user: User = Depends(current_user)) -> dict:
    """What moved since a date: forecast deltas per open question, new ledger rows, new documents, alerts."""
    d = date.fromisoformat(since) if since else date.today() - timedelta(days=1)
    dt = datetime(d.year, d.month, d.day)
    moves = []
    for q in db.scalars(select(Question).where(Question.status == "open").order_by(Question.order)):
        latest = db.scalar(select(Forecast).where(Forecast.question_id == q.id).order_by(Forecast.as_of.desc()).limit(1))
        before = db.scalar(select(Forecast).where(Forecast.question_id == q.id, Forecast.as_of <= d).order_by(Forecast.as_of.desc()).limit(1))
        if latest is None:
            continue
        prev_states = {c["indicator"]: c["state"] for c in ((before.decomposition if before else {}) or {}).get("machine", {}).get("contributions", [])}
        flips = [{"indicator": c["indicator"], "from": prev_states.get(c["indicator"]), "to": c["state"]} for c in (latest.decomposition or {}).get("machine", {}).get("contributions", []) if prev_states.get(c["indicator"]) not in (None, c["state"])]
        moves.append({"code": q.code, "title_en": q.title_en, "title_ar": q.title_ar, "from": before.final_prob if before else None, "to": latest.final_prob, "delta_points": round((latest.final_prob - (before.final_prob if before else latest.final_prob)) * 100, 1), "as_of": latest.as_of.isoformat(), "flips": flips})
    moves.sort(key=lambda m: abs(m["delta_points"]), reverse=True)
    new_events = db.scalars(select(ThreatEvent).where(ThreatEvent.created_at >= dt, ThreatEvent.is_synthetic.is_(False)).order_by(ThreatEvent.severity.desc(), ThreatEvent.event_date.desc()).limit(50)).all()
    new_obs = db.scalar(select(func.count(PolicyObservation.id)).where(PolicyObservation.created_at >= dt)) if hasattr(PolicyObservation, "created_at") else 0
    new_docs = db.scalar(select(func.count(Document.id)).where(Document.fetched_at >= dt))
    alerts = db.scalars(select(Alert).where(Alert.created_at >= dt).order_by(Alert.severity.desc(), Alert.created_at.desc()).limit(50)).all()
    return {
        "since": d.isoformat(), "moves": moves,
        "events": [{"id": e.id, "event_date": e.event_date.isoformat(), "event_type": e.event_type, "country": e.country, "severity": e.severity, "description_en": e.description_en, "description_ar": e.description_ar, "verification": e.verification, "source": (e.provenance or {}).get("source")} for e in new_events],
        "new_observations": new_obs, "new_documents": new_docs, "alerts": [_alert_row(a) for a in alerts],
    }


# ---------------------------------------------------------------------------
# Search and document routing
# ---------------------------------------------------------------------------


def _doc_row(d: Document, snippet_for: str | None = None) -> dict:
    text = d.raw_text or ""
    snippet = text[:280]
    if snippet_for:
        i = text.lower().find(snippet_for.lower())
        if i >= 0:
            snippet = text[max(0, i - 120): i + 200]
    return {"id": d.id, "source": d.source.code, "source_name": d.source.name, "external_id": d.external_id, "title": d.title, "url": d.url, "language": d.language, "published_at": d.published_at.isoformat() if d.published_at else None, "fetched_at": d.fetched_at.isoformat(timespec="seconds"), "status": d.status, "meta": d.meta or {}, "snippet": " ".join(snippet.split())}


@router.get("/search")
def search(q: str = Query(min_length=2, max_length=120), days: int = 365, kind: str = "all", limit: int = 40, db: Session = Depends(get_session), user: User = Depends(current_user)) -> dict:
    since = date.today() - timedelta(days=days)
    like = f"%{q}%"
    out: dict = {"q": q, "events": [], "observations": [], "documents": []}
    if kind in ("all", "events"):
        for e in db.scalars(select(ThreatEvent).where(ThreatEvent.event_date >= since, or_(ThreatEvent.description_en.ilike(like), ThreatEvent.description_ar.ilike(like), ThreatEvent.location_name.ilike(like))).order_by(ThreatEvent.event_date.desc()).limit(limit)):
            out["events"].append({"id": e.id, "event_date": e.event_date.isoformat(), "event_type": e.event_type, "country": e.country, "severity": e.severity, "description_en": e.description_en, "description_ar": e.description_ar, "verification": e.verification, "provenance": e.provenance, "is_synthetic": e.is_synthetic})
    if kind in ("all", "observations"):
        for o in db.scalars(select(PolicyObservation).where(PolicyObservation.observed_at >= since, or_(PolicyObservation.text_excerpt.ilike(like), PolicyObservation.text_excerpt_en.ilike(like))).order_by(PolicyObservation.observed_at.desc()).limit(limit)):
            out["observations"].append({"id": o.id, "actor": o.actor.code, "issue": o.issue.code, "observed_at": o.observed_at.isoformat(), "track": o.track, "score": o.score, "text_excerpt": o.text_excerpt, "text_excerpt_en": o.text_excerpt_en, "verification": o.verification, "provenance": o.provenance})
    if kind in ("all", "documents"):
        sdt = datetime(since.year, since.month, since.day)
        for d in db.scalars(select(Document).where(Document.fetched_at >= sdt, or_(Document.title.ilike(like), Document.raw_text.ilike(like))).order_by(Document.fetched_at.desc()).limit(limit)):
            out["documents"].append(_doc_row(d, q))
    auth.audit(db, user, "search", q[:80], {"kind": kind})
    return out


@router.get("/documents")
def list_documents(source: str | None = None, days: int = 14, limit: int = 100, db: Session = Depends(get_session), user: User = Depends(current_user)) -> list[dict]:
    sdt = datetime.utcnow() - timedelta(days=days)
    stmt = select(Document).where(Document.fetched_at >= sdt).order_by(Document.fetched_at.desc()).limit(limit)
    if source:
        src = db.scalar(select(Source).where(Source.code == source))
        if src is None:
            return []
        stmt = stmt.where(Document.source_id == src.id)
    return [_doc_row(d) for d in db.scalars(stmt)]


@router.get("/documents/{doc_id}")
def document(doc_id: int, db: Session = Depends(get_session), user: User = Depends(current_user)) -> dict:
    d = db.get(Document, doc_id)
    if d is None:
        raise HTTPException(404, "document not found")
    row = _doc_row(d)
    row["text"] = d.raw_text or ""
    row["events"] = [{"id": e.id, "event_type": e.event_type, "event_date": e.event_date.isoformat(), "verification": e.verification, "description_en": e.description_en} for e in db.scalars(select(ThreatEvent).where(ThreatEvent.source_doc_id == d.id))]
    row["observations"] = [{"id": o.id, "actor": o.actor.code, "issue": o.issue.code, "score": o.score, "verification": o.verification} for o in db.scalars(select(PolicyObservation).where(PolicyObservation.source_doc_id == d.id))] if hasattr(PolicyObservation, "source_doc_id") else []
    return row


def routing_terms(q: Question) -> list[str]:
    """Terms that identify material relevant to a question: distinctive words in its title plus
    the actors, countries and event types its indicators watch."""
    terms = {w.lower() for w in re.findall(r"[A-Za-z][A-Za-z\-']{4,}", q.title_en) if w.lower() not in STOP}
    names = {"IR": "iran", "SA": "saudi", "AE": "uae", "QA": "qatar", "KW": "kuwait", "BH": "bahrain", "OM": "oman", "IQ": "iraq", "YE": "yemen", "IL": "israel", "US": "united states", "UK": "united kingdom", "HORMUZ": "hormuz", "RED_SEA": "red sea", "GULF_OF_ADEN": "aden", "YE-HOUTHI": "houthi", "GCC": "gcc"}
    for ind in q.indicators:
        rule = ind.rule or {}
        for c in rule.get("countries", []) + ([rule["actor"]] if rule.get("actor") else []) + rule.get("actors", []):
            if c in names:
                terms.add(names[c])
        for t in rule.get("event_types", []):
            terms.add(t.split("_")[0])
    return sorted(terms)


@router.get("/questions/{code}/documents")
def question_documents(code: str, days: int = 14, limit: int = 40, db: Session = Depends(get_session), user: User = Depends(current_user)) -> dict:
    q = db.scalar(select(Question).where(Question.code == code))
    if q is None:
        raise HTTPException(404, "question not found")
    terms = routing_terms(q)
    sdt = datetime.utcnow() - timedelta(days=days)
    scored: list[tuple[int, Document]] = []
    for d in db.scalars(select(Document).where(Document.fetched_at >= sdt).order_by(Document.fetched_at.desc()).limit(3000)):
        hits = 0
        if code in ((d.meta or {}).get("question_codes") or []):
            hits += 3
        text = ((d.title or "") + " " + (d.raw_text or "")[:4000]).lower()
        hits += sum(1 for t in terms if t in text)
        if hits >= 2:
            scored.append((hits, d))
    scored.sort(key=lambda x: (x[0], x[1].fetched_at), reverse=True)
    return {"question": code, "terms": terms, "documents": [dict(_doc_row(d), score=h) for h, d in scored[:limit]]}


# ---------------------------------------------------------------------------
# Brief export (Word)
# ---------------------------------------------------------------------------


def markdown_to_docx(md: str, *, lang: str, marking: str, title_fallback: str) -> bytes:
    from docx import Document as Docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    doc = Docx()
    rtl = lang == "ar"
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(11)

    def bidi(p) -> None:
        if not rtl:
            return
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        pPr = p._p.get_or_add_pPr()
        el = OxmlElement("w:bidi")
        el.set(qn("w:val"), "1")
        pPr.append(el)

    for section in doc.sections:
        hp = section.header.paragraphs[0]
        hp.text = marking
        hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fp = section.footer.paragraphs[0]
        fp.text = f"{marking}   |   Rasad   |   {datetime.utcnow():%Y-%m-%d %H:%M} UTC"
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    had_title = False
    bullets: list[str] = []

    def flush() -> None:
        nonlocal bullets
        for b in bullets:
            p = doc.add_paragraph(b.strip(), style="List Bullet")
            bidi(p)
        bullets = []

    for ln in md.split("\n"):
        if ln.startswith("# "):
            flush()
            bidi(doc.add_heading(ln[2:].strip(), level=1))
            had_title = True
        elif ln.startswith("## "):
            flush()
            bidi(doc.add_heading(ln[3:].strip(), level=2))
        elif ln.startswith("- ") or ln.startswith("    - "):
            bullets.append(ln.replace("    - ", "- ")[2:])
        elif ln.strip() == "":
            flush()
        else:
            flush()
            text = ln.strip()
            p = doc.add_paragraph()
            if text.startswith("_") and text.endswith("_") and len(text) > 2:
                r = p.add_run(text[1:-1])
                r.italic = True
            else:
                p.add_run(text)
            bidi(p)
    flush()
    if not had_title:
        doc.paragraphs[0].insert_paragraph_before(title_fallback, style="Heading 1")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@router.get("/guide")
def guide(lang: str = Query("en", pattern="^(en|ar)$"), user: User = Depends(current_user)) -> dict:
    """The analyst guide shipped with this version (Arabic falls back to English until translated)."""
    base = os.path.join(os.path.dirname(os.path.dirname(__file__)), "docs")
    check_language(lang)
    path = os.path.join(base, f"guide_{lang}.md")
    if not os.path.exists(path):
        path, lang = os.path.join(base, "guide_en.md"), "en"
    with open(path, encoding="utf-8") as f:
        return {"lang": lang, "content_md": f.read()}


@router.get("/briefs/weekly")
def weekly_brief(lang: str = Query("en", pattern="^(en|ar)$"), as_of: str | None = None, days: int = 7, db: Session = Depends(get_session), user: User = Depends(current_user)) -> dict:
    check_language(lang)
    from ..briefs import weekly

    d = date.fromisoformat(as_of) if as_of else date.today()
    b = weekly.generate(db, d, lang, days)
    auth.audit(db, user, "weekly_brief_generated", d.isoformat(), {"lang": lang})
    return {"id": b.id, "lang": b.lang, "generator": b.generator, "generated_at": b.generated_at.isoformat(), "content_md": b.content_md}


@router.get("/briefs/weekly/docx")
def weekly_brief_docx(lang: str = Query("en", pattern="^(en|ar)$"), as_of: str | None = None, days: int = 7, db: Session = Depends(get_session), user: User = Depends(current_user)):
    check_language(lang)
    from ..briefs import weekly

    d = date.fromisoformat(as_of) if as_of else date.today()
    md = weekly.render(db, d, lang, days)
    marking = os.getenv("RASAD_MARKING", "OFFICIAL")
    data = markdown_to_docx(md, lang=lang, marking=marking, title_fallback="Weekly situation brief")
    auth.audit(db, user, "weekly_brief_exported", d.isoformat(), {"lang": lang})
    return StreamingResponse(io.BytesIO(data), media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document", headers={"Content-Disposition": f'attachment; filename="rasad-weekly-{d.isoformat()}-{lang}.docx"'})


@router.get("/briefs/{code}/docx")
def brief_docx(code: str, lang: str = Query("en", pattern="^(en|ar)$"), as_of: str | None = None, polish: bool = True, db: Session = Depends(get_session), user: User = Depends(current_user)):
    check_language(lang)
    d = date.fromisoformat(as_of) if as_of else date.today()
    try:
        b = briefs.generate(db, code, d, lang, polish=polish)
    except KeyError:
        raise HTTPException(404, "question not found")
    marking = os.getenv("RASAD_MARKING", "OFFICIAL")
    data = markdown_to_docx(b.content_md, lang=lang, marking=marking, title_fallback=code)
    auth.audit(db, user, "brief_exported", code, {"lang": lang, "format": "docx"})
    fname = f"rasad-{code}-{d.isoformat()}-{lang}.docx"
    return StreamingResponse(io.BytesIO(data), media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document", headers={"Content-Disposition": f'attachment; filename="{fname}"'})
